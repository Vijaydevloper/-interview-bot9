import datetime as dt
import json
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import llm
import progress_tools as p
import rag
import scheduler
import ui


def make_corpus(root: Path) -> str:
    (root / "sql").mkdir(parents=True)
    (root / "ml").mkdir()
    (root / "sql" / "joins.md").write_text(
        "# SQL joins\n\nAn INNER JOIN returns rows with matching keys. A LEFT JOIN keeps every row "
        "from the left table and fills NULLs where the right table has no match.\n\n"
        "## Window functions\n\nROW_NUMBER, RANK and LAG operate over a window defined by OVER.\n"
    )
    (root / "ml" / "bias.md").write_text(
        "# Bias and variance\n\nOverfitting means low training error but high validation error. "
        "Regularisation and more data reduce variance.\n\n"
        "# Metrics\n\nPrecision is TP over TP plus FP. Recall is TP over TP plus FN.\n"
    )
    from docx import Document

    d = Document()
    d.add_heading("Python containers", 1)
    d.add_paragraph("Lists are mutable and tuples are immutable.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Type", "Mutable"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "list", "yes"
    (root / "python").mkdir()
    d.save(str(root / "python" / "containers.docx"))
    return str(root)


# ---------- progress ----------
def test_progress_persists_and_weak_topics(tmp_path, monkeypatch):
    monkeypatch.setenv("PROGRESS_DB", str(tmp_path / "t.db"))
    p.log_answer("sql", 3, 10, question="q1", feedback="f")
    p.log_answer("sql", 5, 10)
    p.log_answer("python", 9, 10)
    p.log_answer("ml", 4, 9)  # normalised to /10
    weak = json.loads(p.get_weak_topics())
    assert [w["topic"] for w in weak] == ["sql", "ml"]
    assert weak[0]["attempt_count"] == 2 and weak[0]["average_score"] == 4.0
    assert len(json.loads(p.get_history())) == 4


def test_log_answer_validates(tmp_path, monkeypatch):
    monkeypatch.setenv("PROGRESS_DB", str(tmp_path / "t.db"))
    for bad in [dict(topic="", score=1), dict(topic="x", score=11), dict(topic="x", score=-1)]:
        try:
            p.log_answer(**bad)
            assert False, bad
        except ValueError:
            pass


# ---------- RAG ----------
def test_rag_retrieval_sections_and_docx_tables(tmp_path):
    idx = rag.Index.build(make_corpus(tmp_path / "docs"))
    assert set(idx.topics) == {"sql", "ml", "python"}
    hit = idx.search("difference between LEFT JOIN and INNER JOIN", k=1)[0]
    assert hit.topic == "sql" and hit.score > 0.12
    assert all(c.topic == "ml" for c in idx.search("precision recall", k=3, topic="ml"))
    assert idx.search("zzzz qqqq") == []
    py = idx.search("list tuple mutable", k=1)[0]
    assert py.topic == "python" and "list | yes" in py.text  # table row survived
    assert idx.search("overfitting", k=1)[0].section  # heading carried into the chunk


def test_chunking_overlap_and_limits():
    text = "\n\n".join(f"paragraph {i} " + "word " * 60 for i in range(20))
    chunks = rag.chunk_text(text, size=500, overlap=50)
    assert len(chunks) > 3 and all(len(c) <= 1100 for c in chunks)


def test_heading_stays_with_next_paragraph():
    chunks = rag.chunk_text("# Title\n\nbody text", size=900)
    assert chunks == ["# Title\nbody text"]


# ---------- evaluation ----------
def test_rubric_caps_enforced():
    ev = llm._finalise({"correctness": 9, "coverage": 9, "clarity": 9, "example": 9,
                        "example_applicable": False}, "llm")
    assert ev["score"] == 9 and ev["max_score"] == 9
    ev = llm._finalise({"correctness": 9, "coverage": 9, "clarity": 9, "example": 9}, "llm")
    assert ev["score"] == 10 and ev["max_score"] == 10


def test_offline_evaluation_runs(tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    idx = rag.Index.build(make_corpus(tmp_path / "docs"))
    chunks = idx.search("overfitting variance", k=2, topic="ml")
    good = llm.evaluate("q", "Overfitting means low training error but high validation error; "
                        "regularisation and more data reduce variance, for example L2.", chunks)
    bad = llm.evaluate("q", "no idea", chunks)
    assert good["score"] > bad["score"]
    assert llm.generate_question(chunks)  # offline question uses the section heading


# ---------- scheduler ----------
def test_scheduler_prefers_weak_and_stale_topics():
    now = dt.datetime(2026, 10, 10, tzinfo=dt.timezone.utc)
    stats = [
        {"topic": "sql", "average_score": 3.0, "last_practiced_at": "2026-10-01T00:00:00+00:00"},
        {"topic": "python", "average_score": 9.0, "last_practiced_at": "2026-10-10T00:00:00+00:00"},
    ]
    w = scheduler.topic_weights(["sql", "python", "ml"], stats, now)
    assert w["sql"] > w["ml"] > w["python"]
    picks = [scheduler.pick_topic(["sql", "python"], stats, random.Random(i)) for i in range(300)]
    assert picks.count("sql") > picks.count("python") * 2


# ---------- ui helpers ----------
def test_scoreboard_escapes_and_skips_inapplicable_example():
    caps = llm.RUBRIC
    html_ = ui.scoreboard(7.5, 9, {"correctness": 3, "coverage": 2, "clarity": 2, "example": 0}, caps)
    assert "7.5" in html_ and "Example" not in html_
    assert "<script>" not in ui.tags(["<script>alert(1)</script>"])
    assert "m-weak" in ui.meter("sql", 3.0) and "m-good" in ui.meter("sql", 8.0)


# ---------- offline question quality / docs whitelist ----------
def test_offline_question_skips_generic_headings():
    A = rag.Chunk(text="x", source="a.docx", topic="sql", section="1. Introduction")
    B = rag.Chunk(text="y", source="a.docx", topic="sql", section="3.2 LEFT JOIN")
    assert llm._offline_question([A, B]) == "Explain in your own words (sql): LEFT JOIN"
    only = llm._offline_question([A])
    assert "sql" in only and only.endswith("x")


def test_fetch_doc_rejects_untrusted_domain_before_network():
    import docs_tools
    assert "trusted" in json.loads(docs_tools.fetch_doc("https://evil.example.com/x"))["error"]
    assert "trusted" in json.loads(docs_tools.fetch_doc("https://pandas.pydata.org.evil.com/"))["error"]


def test_mcp_round_trip_progress(tmp_path, monkeypatch):
    try:
        import mcp  # noqa: F401
    except ImportError:
        return
    monkeypatch.setenv("PROGRESS_DB", str(tmp_path / "m.db"))
    monkeypatch.setenv("USE_MCP", "1")
    import tools
    tools.call("progress", "log_answer", topic="sql", score=4, max_score=10)
    weak = tools.call("progress", "get_weak_topics")
    assert weak[0]["topic"] == "sql" and weak[0]["average_score"] == 4.0


# ---------- LangGraph agent + rag MCP server ----------
from langgraph_stub import use_langgraph_or_stub as _use_langgraph_or_stub  # noqa: E402


def _agent(monkeypatch, tmp_path, fake_docs=None, fake_llm=False):
    _use_langgraph_or_stub()
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.setenv("STUDY_DOCS", make_corpus(tmp_path / "docs"))
    monkeypatch.setenv("PROGRESS_DB", str(tmp_path / "p.db"))
    monkeypatch.setenv("USE_MCP", "1")  # real MCP stdio round trips to the servers
    import agent
    import tools

    agent._graph = None
    calls = []
    real = tools.call

    def spy(server, tool, **args):
        calls.append(f"{server}.{tool}")
        if server == "docs" and fake_docs is not None:
            return fake_docs(tool, args)
        return real(server, tool, **args)

    monkeypatch_attr(agent.tools, "call", spy)
    return agent, calls


_restore = []


def monkeypatch_attr(obj, name, value):
    _restore.append((obj, name, getattr(obj, name)))
    setattr(obj, name, value)


def _undo():
    while _restore:
        o, n, v = _restore.pop()
        setattr(o, n, v)


TOPICS_ = ["ml", "python", "sql"]


def test_rag_mcp_server_tools(tmp_path, monkeypatch):
    monkeypatch.setenv("STUDY_DOCS", make_corpus(tmp_path / "docs"))
    monkeypatch.setenv("USE_MCP", "1")
    import tools

    hits = tools.call("rag", "search_notes", query="LEFT JOIN vs INNER JOIN", k=2)
    assert hits[0]["topic"] == "sql" and hits[0]["score"] > 0.12 and hits[0]["source"] == "joins.md"
    assert tools.call("rag", "list_topics") == {"ml": 1, "python": 1, "sql": 1}
    assert tools.call("rag", "search_notes", query="zzzz qqqq") == []
    s = tools.call("rag", "sample_notes", topic="ml", k=2)
    assert s and all(c["topic"] == "ml" for c in s)
    lib = tools.call("rag", "library_summary")
    assert lib["sql"] == {"chunks": 1, "files": ["joins.md"]} and lib["python"]["files"] == ["containers.docx"]


def test_pptx_slides_are_indexed(tmp_path):
    import zipfile

    root = tmp_path / "docs" / "deep_learning"
    root.mkdir(parents=True)
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    slide = lambda *lines: ("<p:sld xmlns:p='p' %s><p:sp>" % ns + "".join(
        f"<a:p><a:r><a:t>{x}</a:t></a:r></a:p>" for x in lines) + "</p:sp></p:sld>")
    with zipfile.ZipFile(root / "cnn.pptx", "w") as z:
        z.writestr("ppt/slides/slide2.xml", slide("Pooling", "Max pooling keeps the strongest activation."))
        z.writestr("ppt/slides/slide1.xml", slide("Convolution", "A kernel slides over the image."))
    idx = rag.Index.build(str(tmp_path / "docs"))
    hit = idx.search("max pooling activation", k=1)[0]
    assert hit.topic == "deep_learning" and hit.section == "Pooling" and hit.page == 2
    assert idx.search("kernel slides image", k=1)[0].section == "Convolution"


def test_progress_records_mode_and_migrates_old_db(tmp_path, monkeypatch):
    import sqlite3

    db = tmp_path / "old.db"
    c = sqlite3.connect(db)
    c.execute("CREATE TABLE attempts (attempt_id INTEGER PRIMARY KEY AUTOINCREMENT, topic TEXT NOT NULL,"
              " question TEXT NOT NULL DEFAULT '', user_answer TEXT NOT NULL DEFAULT '', score REAL NOT NULL,"
              " max_score REAL NOT NULL DEFAULT 10, feedback TEXT NOT NULL DEFAULT '',"
              " source TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL)")
    c.execute("INSERT INTO attempts (topic, score, created_at) VALUES ('sql', 5, '2026-01-01T00:00:00+00:00')")
    c.commit(); c.close()
    monkeypatch.setenv("PROGRESS_DB", str(db))
    p.log_answer("sql", 8, 10, mode="Coding", difficulty="Advanced")
    rows = json.loads(p.get_history())
    assert rows[0]["mode"] == "Coding" and rows[0]["difficulty"] == "Advanced" and rows[1]["mode"] == ""


def test_agent_question_then_evaluate_logs_progress(tmp_path, monkeypatch):
    try:
        agent, calls = _agent(monkeypatch, tmp_path)
        q = agent.run("question", topic="sql", topics=TOPICS_, difficulty="Beginner", mode="Coding")
        assert q["question"].startswith("Explain in your own words (sql)") and q["context"]
        assert calls == ["rag.sample_notes"]  # no focus -> varied material, not a fixed query
        e = agent.run("evaluate", topic=q["topic"], question=q["question"], context=q["context"],
                      answer_text="An INNER JOIN returns rows with matching keys, for example orders to users.",
                      mode="Coding", difficulty="Beginner")
        last_call = calls[-1]
        import tools
        hist = tools.call("progress", "get_history")
    finally:
        _undo()
    assert last_call == "progress.log_answer" and e["attempt_id"] == 1
    assert e["result"]["mode"] == "offline-heuristic" and 0 <= e["result"]["score"] <= 10
    assert hist[0]["mode"] == "Coding" and hist[0]["difficulty"] == "Beginner" and "joins.md" in hist[0]["source"]


def test_agent_question_with_focus_and_smart_pick(tmp_path, monkeypatch):
    try:
        agent, calls = _agent(monkeypatch, tmp_path)
        f = agent.run("question", topic="sql", topics=TOPICS_, focus="left join")
        assert calls == ["rag.search_notes"] and f["context"][0]["topic"] == "sql"
        calls.clear()
        s = agent.run("question", topic="Smart pick", topics=TOPICS_)
    finally:
        _undo()
    assert calls[0] == "progress.get_topic_stats" and s["topic"] in TOPICS_ and s["question"]


def test_agent_revision_notes(tmp_path, monkeypatch):
    try:
        agent, calls = _agent(monkeypatch, tmp_path)
        import tools
        tools.call("progress", "log_answer", topic="ml", score=3, max_score=10, question="q", feedback="weak")
        r = agent.run("revise", topic="ml", focus="overfitting")
    finally:
        _undo()
    assert calls[-2:] == ["progress.get_topic_stats", "progress.get_history"]
    assert "ml (current average 3.0/10)" in r["sheet"] and r["sources"]


def test_agent_explain_from_notes_without_web(tmp_path, monkeypatch):
    try:
        agent, calls = _agent(monkeypatch, tmp_path)
        out = agent.run("explain", question="What is the difference between LEFT JOIN and INNER JOIN?",
                        topic=None, use_web=True)
    finally:
        _undo()
    assert out["covered"] and calls == ["rag.search_notes"]  # notes were enough: docs never called
    assert any(s["kind"] == "notes" and "joins.md" in s["label"] for s in out["sources"])


def test_agent_explain_falls_back_to_official_docs(tmp_path, monkeypatch):
    def fake_docs(tool, args):
        if tool == "search_docs":
            return [{"title": "t", "url": "https://docs.python.org/3/x", "snippet": ""}]
        return {"url": args["url"], "text": "Generators yield values lazily."}

    try:
        agent, calls = _agent(monkeypatch, tmp_path, fake_docs)
        out = agent.run("explain", question="How do generators yield values lazily?", topic=None, use_web=True)
    finally:
        _undo()
    assert not out["covered"]
    assert calls == ["rag.search_notes", "docs.search_docs", "docs.fetch_doc"]
    assert out["sources"] == [{"kind": "docs", "label": "https://docs.python.org/3/x"}]
    assert "Generators" in out["answer"]


def test_agent_explain_declines_and_reports_failures(tmp_path, monkeypatch):
    try:
        agent, calls = _agent(monkeypatch, tmp_path)
        out = agent.run("explain", question="explain quantum chromodynamics", topic=None, use_web=False)
        assert calls == ["rag.search_notes"] and out["sources"] == [] and "won't guess" in out["answer"]
        _undo()
        agent, _ = _agent(monkeypatch, tmp_path / "second",
                          lambda tool, args: {"error": "search failed: offline"})
        out = agent.run("explain", question="explain quantum chromodynamics", topic=None, use_web=True)
    finally:
        _undo()
    assert "Docs lookup failed: search failed: offline" in out["answer"]



# ---------- Groq backend (mock OpenAI-compatible server) ----------
def _mock_groq(script):
    """script: list of (status, headers, content-or-error-text); returns (server, seen requests)."""
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    seen = []

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            seen.append((self.path, self.headers.get("Authorization"), body))
            status, headers, content = script.pop(0)
            out = json.dumps({"choices": [{"message": {"content": content}}]} if status == 200
                             else {"error": {"message": content}}).encode()
            self.send_response(status)
            for k, v in headers.items():
                self.send_header(k, v)
            self.send_header("Content-Length", str(len(out)))
            self.end_headers()
            self.wfile.write(out)

        def log_message(self, *a):
            pass

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, seen


def _groq_env(monkeypatch, srv):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_test")
    monkeypatch.setenv("GROQ_BASE_URL", f"http://127.0.0.1:{srv.server_port}/openai/v1")
    monkeypatch.setenv("GROQ_MODEL", "tiny-model")


def test_groq_question_and_evaluation(tmp_path, monkeypatch):
    good = {"correctness": 3, "coverage": 2, "clarity": 2, "example": 1,
            "correct_points": ["a"], "missing_points": "b", "model_answer": "m", "feedback": "ok"}
    srv, seen = _mock_groq([(200, {}, "What is a LEFT JOIN?"), (200, {}, "not json"),
                            (200, {}, json.dumps(good))])
    try:
        _groq_env(monkeypatch, srv)
        idx = rag.Index.build(make_corpus(tmp_path / "docs"))
        chunks = idx.search("left join", k=2, topic="sql")
        assert llm.online()
        assert llm.generate_question(chunks, difficulty="Advanced", mode="Coding") == "What is a LEFT JOIN?"
        ev = llm.evaluate("q", "a", chunks)
    finally:
        srv.shutdown()
    assert ev["mode"] == "llm" and ev["score"] == 8 and ev["missing_points"] == ["b"]
    path, auth, body = seen[0]
    assert path == "/openai/v1/chat/completions" and auth == "Bearer gsk_test" and body["model"] == "tiny-model"
    assert "advanced" in body["messages"][1]["content"] and "coding" in body["messages"][1]["content"]
    assert seen[1][2]["response_format"] == {"type": "json_object"} and len(seen) == 3  # retried bad JSON


def test_groq_rate_limit_retry_and_auth_error(monkeypatch):
    srv, seen = _mock_groq([(429, {"retry-after": "0"}, "slow down"), (200, {}, "hello"),
                            (401, {}, "bad key")])
    try:
        _groq_env(monkeypatch, srv)
        assert llm._complete("s", "u") == "hello" and len(seen) == 2
        try:
            llm._complete("s", "u")
            assert False
        except RuntimeError as e:
            assert "401" in str(e) and "GROQ_API_KEY" in str(e)
    finally:
        srv.shutdown()


def test_groq_unreachable_and_offline(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_test")
    monkeypatch.setenv("GROQ_BASE_URL", "http://127.0.0.1:1/openai/v1")
    try:
        llm._complete("s", "u")
        assert False
    except RuntimeError as e:
        assert "Cannot reach Groq" in str(e)
    monkeypatch.delenv("GROQ_API_KEY")
    assert not llm.online()


# ---------- .env loading ----------
def test_env_file_loaded_without_overriding(tmp_path, monkeypatch):
    import config

    f = tmp_path / ".env"
    f.write_text("# comment\nGROQ_API_KEY=gsk_from_file\nGROQ_MODEL='m1'\nexport EXTRA=1\n")
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.setenv("GROQ_MODEL", "already-set")
    monkeypatch.delenv("EXTRA", raising=False)
    assert config.load_env(f) is True
    import os

    assert os.environ["GROQ_API_KEY"] == "gsk_from_file" and os.environ["EXTRA"] == "1"
    assert os.environ["GROQ_MODEL"] == "already-set"  # real environment wins
    monkeypatch.delenv("GROQ_API_KEY"); monkeypatch.delenv("EXTRA")
    assert config.load_env(tmp_path / "missing.env") is False


# ---------- MCQ ----------
def test_mcq_bank_has_ten_valid_questions_and_shuffle_keeps_answers():
    import mcq

    bank = mcq.load_bank()
    assert len(bank) == 10
    for q in bank:
        assert len(q["options"]) == 4 and 0 <= q["answer"] < 4 and q["explanation"] and q["topic"]
    right = {q["question"]: q["options"][q["answer"]] for q in bank}
    for seed in (1, 2, 3):
        shuffled = mcq.prepare(bank, seed)
        assert len(shuffled) == 10
        for q in shuffled:
            assert q["options"][q["answer"]] == right[q["question"]]
    g = mcq.grade(shuffled, [q["answer"] for q in shuffled[:6]] + [(q["answer"] + 1) % 4 for q in shuffled[6:]])
    assert g["correct"] == 6 and g["total"] == 10


def test_mcq_validate_rejects_bad_items():
    import mcq

    ok = {"question": "q?", "options": ["a", "b", "c", "d"], "answer": 2}
    assert mcq.validate(ok)["topic"] == "general"
    for bad in [{**ok, "options": ["a", "a", "c", "d"]}, {**ok, "answer": 4}, {**ok, "options": ["a", "b"]},
                {**ok, "question": " "}, {"question": "q"}]:
        assert mcq.validate(bad) is None


def test_log_answers_batch_is_all_or_nothing(tmp_path, monkeypatch):
    monkeypatch.setenv("PROGRESS_DB", str(tmp_path / "t.db"))
    ids = json.loads(p.log_answers([{"topic": "sql", "score": 10}, {"topic": "ml", "score": 0, "mode": "MCQ"}]))
    assert len(ids["attempt_ids"]) == 2
    try:
        p.log_answers([{"topic": "sql", "score": 5}, {"topic": "", "score": 1}])
        raise AssertionError("expected ValueError")
    except ValueError:
        pass
    assert len(json.loads(p.get_history())) == 2  # the bad batch wrote nothing


def test_groq_generates_mcqs_and_chat(tmp_path, monkeypatch):
    docs = make_corpus(tmp_path / "docs")
    chunks = rag.Index.build(docs).search("join", k=2)
    good = {"topic": "sql", "question": "What does LEFT JOIN keep?", "options": ["left rows", "right rows", "none", "all"],
            "answer": 0, "explanation": "left", "source": "joins.md"}
    body = json.dumps({"questions": [good, {"question": "broken"}]})
    srv, seen = _mock_groq([(200, {}, body), (200, {}, "Hi! **Q: what is a JOIN?**"), (200, {}, "Good. Next: ...")])
    try:
        _groq_env(monkeypatch, srv)
        qs = llm.generate_mcqs(chunks, 10)
        assert len(qs) == 1 and qs[0]["answer"] == 0  # the malformed item was dropped
        first = llm.chat_reply([], chunks, "Sql", "Beginner")
        assert "JOIN" in first
        second = llm.chat_reply([{"role": "assistant", "content": first}, {"role": "user", "content": "it joins tables"}],
                                chunks, "Sql", "Beginner")
        assert second.startswith("Good")
        sent = seen[-1][2]["messages"]  # system prompt carries the notes; the history follows it
        assert sent[0]["role"] == "system" and "INNER JOIN" in sent[0]["content"]
        assert [m["role"] for m in sent[1:]] == ["assistant", "user"]
    finally:
        srv.shutdown()


def test_chat_offline_asks_then_checks(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    chunks = [rag.Chunk("An INNER JOIN returns matching rows. LEFT JOIN keeps left rows.", "sql.md", "sql", None, 0, "SQL joins")]
    q = llm.chat_reply([], chunks)
    assert "Offline" in q and "SQL joins" in q
    r = llm.chat_reply([{"role": "assistant", "content": q}, {"role": "user", "content": "a join returns matching rows"}], chunks)
    assert "Next question" in r and "/" in r
